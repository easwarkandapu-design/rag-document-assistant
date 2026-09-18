import os
import re
import time
from pathlib import Path

from pypdf import PdfReader
from docx import Document as DocxDocument
from google import genai
from google.genai import types
from sqlalchemy import text

from .extensions import db
from .models import Document, DocumentChunk


client = None


# =========================================================
# GEMINI CLIENT
# =========================================================

def get_client():
    global client

    if client is None:
        api_key = os.getenv("GEMINI_API_KEY")

        if not api_key:
            raise RuntimeError(
                "GEMINI_API_KEY is not configured in .env"
            )

        client = genai.Client(api_key=api_key)

    return client


# =========================================================
# EXTRACT TEXT
# =========================================================

def extract_text(path: str) -> str:

    suffix = Path(path).suffix.lower()

    if suffix == ".pdf":
        reader = PdfReader(path)

        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    if suffix == ".docx":
        doc = DocxDocument(path)

        return "\n".join(
            paragraph.text
            for paragraph in doc.paragraphs
        )

    if suffix in {".txt", ".md"}:
        return Path(path).read_text(
            encoding="utf-8",
            errors="ignore"
        )

    raise ValueError(
        "Unsupported file type. Use PDF, DOCX, TXT, or MD."
    )


# =========================================================
# NORMALIZE TEXT
# =========================================================

def normalize_text(text_value: str) -> str:

    text_value = text_value.replace(
        "\x00",
        " "
    )

    text_value = re.sub(
        r"[ \t]+",
        " ",
        text_value
    )

    text_value = re.sub(
        r"\n{3,}",
        "\n\n",
        text_value
    )

    return text_value.strip()


# =========================================================
# CHUNK DOCUMENT
# =========================================================

def chunk_text(text_value: str):

    size = int(
        os.getenv(
            "CHUNK_SIZE",
            "900"
        )
    )

    overlap = int(
        os.getenv(
            "CHUNK_OVERLAP",
            "150"
        )
    )

    if overlap >= size:
        raise ValueError(
            "CHUNK_OVERLAP must be smaller than CHUNK_SIZE"
        )

    chunks = []

    start = 0
    length = len(text_value)

    while start < length:

        end = min(
            start + size,
            length
        )

        if end < length:

            paragraph_boundary = text_value.rfind(
                "\n\n",
                start,
                end
            )

            sentence_boundary = text_value.rfind(
                ". ",
                start,
                end
            )

            boundary = max(
                paragraph_boundary,
                sentence_boundary
            )

            if boundary > start + size // 2:
                end = boundary + 2

        chunk = text_value[
            start:end
        ].strip()

        if chunk:
            chunks.append(chunk)

        if end >= length:
            break

        start = max(
            end - overlap,
            start + 1
        )

    return chunks


# =========================================================
# EMBEDDINGS
# =========================================================

def embed_texts(texts):

    response = get_client().models.embed_content(

        model=os.getenv(
            "GEMINI_EMBED_MODEL",
            "gemini-embedding-001"
        ),

        contents=texts,

        config=types.EmbedContentConfig(
            output_dimensionality=3072
        )
    )

    return [
        item.values
        for item in response.embeddings
    ]


# =========================================================
# INGEST DOCUMENT
# =========================================================

def ingest_document(path: str, filename: str):

    raw = normalize_text(
        extract_text(path)
    )

    if not raw:
        raise ValueError(
            "No readable text was found in the document."
        )

    chunks = chunk_text(raw)

    embeddings = embed_texts(chunks)

    document = Document(
        filename=filename,
        file_type=Path(
            filename
        ).suffix.lower()
    )

    db.session.add(document)

    db.session.flush()

    for index, (chunk, vector) in enumerate(
        zip(chunks, embeddings)
    ):

        db.session.add(
            DocumentChunk(
                document_id=document.id,
                chunk_index=index,
                content=chunk,
                embedding=vector
            )
        )

    db.session.commit()

    return document.id, len(chunks)


# =========================================================
# RETRIEVE FROM POSTGRESQL
# =========================================================

def retrieve(question: str, top_k=None):

    top_k = top_k or int(
        os.getenv(
            "TOP_K",
            "5"
        )
    )

    # Convert user question into an embedding
    question_vector = embed_texts(
        [question]
    )[0]

    # IMPORTANT:
    # Retrieval happens ONLY from PostgreSQL / pgvector.
    rows = db.session.execute(

        text(
            """
            SELECT
                dc.id,
                dc.document_id,
                dc.chunk_index,
                dc.content,
                d.filename,

                1 - (
                    dc.embedding
                    <=>
                    CAST(:qvec AS vector)
                ) AS similarity

            FROM document_chunks dc

            JOIN documents d
                ON d.id = dc.document_id

            ORDER BY
                dc.embedding
                <=>
                CAST(:qvec AS vector)

            LIMIT :top_k
            """
        ),

        {
            "qvec": str(question_vector),
            "top_k": top_k
        }

    ).mappings().all()

    unique_rows = []

    seen = set()

    for row in rows:

        key = (
            row["document_id"],
            row["chunk_index"]
        )

        if key in seen:
            continue

        seen.add(key)

        unique_rows.append(
            dict(row)
        )

    return unique_rows


# =========================================================
# BUILD EVIDENCE
# =========================================================

def build_evidence(rows):

    parts = []

    for number, row in enumerate(
        rows,
        start=1
    ):

        parts.append(
            f"""
[SOURCE {number} | {row['filename']} | chunk {row['chunk_index']}]

{row['content']}
""".strip()
        )

    return "\n\n".join(parts)


# =========================================================
# GEMINI GENERATION WITH RETRY
# =========================================================

def generate_with_retry(
    model,
    prompt,
    attempts=4
):

    last_error = None

    for attempt in range(attempts):

        try:

            response = get_client().models.generate_content(

                model=model,

                contents=prompt,

                config=types.GenerateContentConfig(
                    temperature=0
                )
            )

            answer = ""

            if response is not None:
                answer = (
                    getattr(response, "text", "") or ""
                ).strip()

            if answer:
                return answer

            raise RuntimeError(
                "Gemini returned an empty response."
            )

        except Exception as error:

            last_error = error

            print(
                f"[GEMINI ERROR] "
                f"Attempt {attempt + 1}/{attempts}: "
                f"{type(error).__name__}: {error}"
            )

            if attempt < attempts - 1:

                # Retry all temporary/API failures.
                # This also helps with transient Gemini failures.
                wait_time = 2 ** attempt

                print(
                    f"[GEMINI] Retrying in "
                    f"{wait_time} seconds..."
                )

                time.sleep(wait_time)

    raise last_error


# =========================================================
# GENERATE ANSWER
# =========================================================

def generate_answer(
    question,
    rows,
    strict=False
):

    evidence = build_evidence(rows)

    if strict:

        rules = """
You are a strict document-grounded RAG assistant.

Answer ONLY using the PostgreSQL evidence provided below.

Rules:
- Do not use outside knowledge.
- Do not guess.
- Do not invent information.
- Every factual statement must be supported by the evidence.
- Cite the supporting evidence using [SOURCE 1], [SOURCE 2], etc.
- Use only source numbers that actually exist.
- If the evidence does not answer the question, say exactly:

The uploaded documents do not provide enough information to answer that.
"""

    else:

        rules = """
You are a document-grounded RAG assistant.

Answer ONLY using the PostgreSQL evidence provided below.

Rules:
- Do not use outside knowledge.
- Do not guess.
- Do not invent information.
- Give a clear and useful answer.
- Cite the supporting evidence using [SOURCE 1], [SOURCE 2], etc.
- Use only source numbers that actually exist.
- If the evidence does not answer the question, say exactly:

The uploaded documents do not provide enough information to answer that.
"""

    prompt = f"""
{rules}

USER QUESTION:
{question}

POSTGRESQL EVIDENCE:
{evidence}

Now answer the user's question.
"""

    model = os.getenv(
        "GEMINI_CHAT_MODEL",
        "gemini-3.6-flash-lite"
    )

    return generate_with_retry(
        model,
        prompt
    )


# =========================================================
# VALIDATOR
# =========================================================

def validate_answer(
    answer,
    rows
):

    if not answer:
        return False, "empty"

    answer = answer.strip()

    if not rows:
        return False, "no_evidence"

    # Safe fallback is acceptable.
    fallback = (
        "The uploaded documents do not provide enough "
        "information to answer that."
    )

    if fallback.lower() in answer.lower():
        return True, "safe_fallback"

    # -----------------------------------------------------
    # SOURCE CITATION VALIDATION
    # -----------------------------------------------------

    citations = re.findall(
        r"\[SOURCE\s+(\d+)\]",
        answer,
        flags=re.IGNORECASE
    )

    if not citations:
        return False, "missing_source_citations"

    # Make sure cited source numbers exist.
    for citation in citations:

        number = int(citation)

        if number < 1 or number > len(rows):
            return False, "invalid_source_citation"

    return True, "verified"


# =========================================================
# MAIN RAG PIPELINE
# =========================================================

def answer_question(question: str):

    question = (
        question or ""
    ).strip()

    if not question:

        return {
            "answer": "Please enter a question.",
            "sources": [],
            "validated": False,
            "validation_reason": "empty_question"
        }

    try:

        # =================================================
        # STEP 1: RETRIEVE FROM POSTGRESQL
        # =================================================

        rows = retrieve(
            question
        )

        print(
            f"[RAG] Retrieved {len(rows)} chunks"
        )

        for row in rows:

            print(
                f"[RAG] chunk={row['chunk_index']} "
                f"similarity={float(row['similarity']):.4f}"
            )

        if not rows:

            return {
                "answer": (
                    "The uploaded documents do not provide "
                    "enough information to answer that."
                ),
                "sources": [],
                "validated": False,
                "validation_reason": "no_evidence"
            }

        # =================================================
        # STEP 2: GENERATE ANSWER USING GEMINI
        # =================================================

        answer = generate_answer(
            question,
            rows,
            strict=False
        )

        print(
            "[RAG] Gemini answer generated successfully."
        )

        # =================================================
        # STEP 3: VALIDATE ANSWER
        # =================================================

        valid, reason = validate_answer(
            answer,
            rows
        )

        print(
            f"[VALIDATOR] valid={valid}, reason={reason}"
        )

        # =================================================
        # STEP 4: STRICT RETRY
        # =================================================

        if not valid:

            print(
                f"[VALIDATOR] First validation failed: {reason}"
            )

            answer = generate_answer(
                question,
                rows,
                strict=True
            )

            valid, reason = validate_answer(
                answer,
                rows
            )

            print(
                f"[VALIDATOR] Retry valid={valid}, "
                f"reason={reason}"
            )

        # =================================================
        # STEP 5: FINAL SAFE FALLBACK
        # =================================================

        if not valid:

            answer = (
                "The uploaded documents do not provide "
                "enough information to answer that."
            )

            reason = "insufficient_verified_evidence"

        # =================================================
        # STEP 6: SOURCE INFORMATION
        # =================================================

        sources = [

            {
                "filename": row["filename"],
                "chunk_index": row["chunk_index"],
                "similarity": round(
                    float(row["similarity"]),
                    4
                )
            }

            for row in rows
        ]

        # =================================================
        # STEP 7: RETURN RESULT
        # =================================================

        return {
            "answer": answer,
            "sources": sources,
            "validated": valid,
            "validation_reason": reason
        }

    except Exception as error:

        # =================================================
        # IMPORTANT:
        # NEVER HIDE THE REAL ERROR DURING DEVELOPMENT.
        # =================================================

        print("\n")
        print("=" * 60)
        print("[RAG ERROR]")
        print(
            f"Type: {type(error).__name__}"
        )
        print(
            f"Message: {error}"
        )
        print("=" * 60)
        print("\n")

        return {
            "answer": (
                "The RAG pipeline encountered an error. "
                "Please check the Flask terminal for the "
                "exact error message."
            ),
            "sources": [],
            "validated": False,
            "validation_reason": "system_error"
        }